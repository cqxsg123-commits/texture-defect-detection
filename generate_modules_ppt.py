# -*- coding: utf-8 -*-
"""生成《主模型与桌面应用》模块讲解 PPT（2026-10-24 工作总结）。

模块一 主模型：PatchCore 主方案 —— 优化议程、批量前向 + GPU 检索提速、
聚合方式消融、最终指标；
模块二 桌面应用：质检判定系统 —— 自动阈值 + 灵敏度倍数滑杆 + 合格/不合格判定、
PatchCore 留一法、测试集判定验证。

运行：D:\\anaconda\\envs\\defect_det\\python.exe generate_modules_ppt.py
产物：基于纹理特征的工业缺陷检测系统_模块讲解.pptx
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(r"D:\基于纹理特征的工业缺陷检测系统")
DD = ROOT / "defect_detection"
RES = DD / "results"
FIG_DIR = ROOT / "ppt_figs_modules"
FIG_DIR.mkdir(parents=True, exist_ok=True)

# ---------- PPT 设计变量（与实验报告 PPT 同一设计系统） ----------
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)      # 深蓝
ACCENT2 = RGBColor(0x2A, 0x78, 0xD6)     # 蓝
ACCENT_D = RGBColor(0x10, 0x42, 0x81)    # 蓝 700
TINT_B = RGBColor(0xCD, 0xE2, 0xFB)      # 蓝 100 浅底
TINT_B2 = RGBColor(0x9E, 0xC5, 0xF4)     # 蓝 200
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK_C = RGBColor(0x0B, 0x0B, 0x0B)
SEC_C = RGBColor(0x52, 0x51, 0x4E)
MUTED_C = RGBColor(0x89, 0x87, 0x81)
ALT_C = RGBColor(0xF5, 0xF6, 0xF8)
HAIR = RGBColor(0xD9, 0xD9, 0xD9)
BLUE_C = RGBColor(0x2A, 0x78, 0xD6)
ORG_C = RGBColor(0xEB, 0x68, 0x34)
GOOD_C = RGBColor(0x0B, 0x8A, 0x3C)
BAD_C = RGBColor(0xD0, 0x3B, 0x3B)
FONT = "Microsoft YaHei"

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "sans-serif"]
plt.rcParams["axes.unicode_minus"] = False
SURFACE = "#fcfcfb"
INK = "#0b0b0b"

# ================= 模块二演示图（真实推理输出） =================
import sys
sys.path.insert(0, str(DD))
import cv2


def _make_demo_figures():
    """用 app.detect 真实推理输出生成两张演示图（判定图带绿/红标题）。"""
    import app
    app.build_ui()  # 加载三模型与自动阈值
    test_root = DD / "data" / "mvtec_anomaly_detection" / "metal_nut" / "test"

    def strip(path, title, verdict, color, orig, overlay, mask):
        fig, axes = plt.subplots(1, 3, figsize=(8.8, 3.1), dpi=200)
        fig.patch.set_facecolor(SURFACE)
        for ax, im, cap in zip(axes, [orig, overlay[..., ::-1], mask],
                               ["原图", "异常热力图（叠加）", "缺陷掩码（≥ 判定阈值）"]):
            ax.imshow(im, cmap="gray" if im.ndim == 2 else None)
            ax.set_title(cap, fontsize=11.5, color=INK)
            ax.axis("off")
        fig.suptitle(f"{title}　·　判定：{verdict}", fontsize=14,
                     color=color, fontweight="bold", y=1.0)
        fig.tight_layout()
        fig.savefig(path, facecolor=SURFACE, bbox_inches="tight")
        plt.close(fig)
        print("[fig]", path.name)

    # 图 1：PatchCore 主方案 → 缺陷样本（不合格）
    p = sorted(test_root.glob("bent/*.png"))[0]
    img = cv2.imdecode(np.fromfile(str(p), dtype=np.uint8), cv2.IMREAD_COLOR)[..., ::-1]
    orig, overlay, mask, html, txt = app.detect(img, "PatchCore（主方案）", 1.0)
    score = float(txt.splitlines()[0].split("：")[1])
    strip(FIG_DIR / "demo_patchcore_defect.png", "PatchCore（主方案）· 缺陷样本（bent）",
          f"不合格（得分 {score:.1f} ≥ 阈值 1742.6）", "#d03b3b", orig, overlay, mask)

    # 图 2：PaDiM → 正常样本（合格；挑选未误报的正常图）
    auto_p = app._AUTO_THRESHOLDS["PaDiM"]
    for p in sorted(test_root.glob("good/*.png")):
        img = cv2.imdecode(np.fromfile(str(p), dtype=np.uint8), cv2.IMREAD_COLOR)[..., ::-1]
        orig, overlay, mask, html, txt = app.detect(img, "PaDiM", 1.0)
        score = float(txt.splitlines()[0].split("：")[1])
        if "不合格" not in html:
            strip(FIG_DIR / "demo_padim_good.png", "PaDiM · 正常样本（good）",
                  f"合格（得分 {score:.1f} ＜ 阈值 {auto_p:.1f}）", "#0b8a3c", orig, overlay, mask)
            break
    return {"defect": FIG_DIR / "demo_patchcore_defect.png",
            "good": FIG_DIR / "demo_padim_good.png"}


DEMO = _make_demo_figures()

# ================= PPT 基础工具（与实验报告 PPT 一致） =================


def set_run(r, size, bold=False, color=INK_C, font=FONT):
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = font
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        if "space_before" in spec:
            p.space_before = Pt(spec["space_before"])
        if "space_after" in spec:
            p.space_after = Pt(spec["space_after"])
        if "line" in spec:
            p.line_spacing = spec["line"]
        for txt, size, bold, color in spec["runs"]:
            r = p.add_run()
            r.text = txt
            set_run(r, size, bold, color)
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp


def header(slide, num, title, sub=None):
    rect(slide, 0.55, 0.44, 0.16, 0.34, ACCENT2)
    runs = []
    if num:
        runs.append((num + "  ", 20, True, ACCENT2))
    runs.append((title, 20, True, INK_C))
    textbox(slide, 0.88, 0.36, 11.5, 0.55, [{"runs": runs}], anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, 0.55, 1.06, 12.23, 0.022, BLUE_C)
    if sub:
        textbox(slide, 0.55, 1.14, 12.2, 0.3, [{"runs": [(sub, 10.5, False, MUTED_C)]}])


def footer(slide, page):
    textbox(slide, 0.55, 7.14, 6, 0.3,
            [{"runs": [("基于纹理特征的工业缺陷检测系统 · 模块讲解", 8, False, MUTED_C)]}])
    textbox(slide, 11.9, 7.14, 0.88, 0.3,
            [{"runs": [(f"{page:02d}", 9, True, ACCENT2)], "align": PP_ALIGN.RIGHT}])


def make_table(slide, x, y, w, rows, col_w=None, col_aligns=None, header_size=11.5,
               font_size=10.5, row_h=0.42, header_fill=ACCENT,
               bold_cells=None, color_cells=None, fill_cells=None):
    n_r, n_c = len(rows), len(rows[0])
    if col_w is None:
        col_w = [1.0 / n_c] * n_c
    if col_aligns is None:
        col_aligns = [PP_ALIGN.CENTER] * n_c
    bold_cells = bold_cells or {}
    color_cells = color_cells or {}
    fill_cells = fill_cells or {}
    gf = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w),
                                Inches(0.5 + row_h * (n_r - 1)))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    total = sum(col_w)
    for j in range(n_c):
        tbl.columns[j].width = Emu(int(Inches(w) * col_w[j] / total))
    for i in range(n_r):
        tbl.rows[i].height = Inches(0.5 if i == 0 else row_h)
        for j in range(n_c):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
            elif i % 2 == 1:
                cell.fill.fore_color.rgb = WHITE
            else:
                cell.fill.fore_color.rgb = ALT_C
            if (i, j) in fill_cells:
                cell.fill.fore_color.rgb = fill_cells[(i, j)]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = col_aligns[j]
            r = p.add_run()
            r.text = str(rows[i][j])
            if i == 0:
                set_run(r, header_size, True, WHITE)
            else:
                set_run(r, font_size, bold_cells.get((i, j), False),
                        color_cells.get((i, j), SEC_C))
    return tbl


def img_w_for_h(path: Path, h: float) -> float:
    w0, h0 = PILImage.open(path).size
    return w0 * h / h0


def pic(slide, path, x, y, h):
    w = img_w_for_h(path, h)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    return w


def bullets(specs, size=12, color=INK_C, line=1.25, gap=8, align=PP_ALIGN.LEFT):
    paras = []
    for txt, *rest in specs:
        bold = rest[0] if rest else False
        paras.append({
            "runs": [("•  ", size, True, ACCENT2), (txt, size, bold, color)],
            "line": line,
            "space_after": gap,
            "align": align,
        })
    return paras


# ================= 构建 PPT =================
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
prs.core_properties.title = "基于纹理特征的工业缺陷检测系统 模块讲解"
prs.core_properties.author = "陈宇 202305630103"

PAGE = 0


def new_slide():
    global PAGE
    PAGE += 1
    return prs.slides.add_slide(BLANK)


def content_slide(num, title, sub=None):
    s = new_slide()
    header(s, num, title, sub)
    return s


def divider(num, big, sub, points):
    s = new_slide()
    rect(s, 0, 0, 13.333, 7.5, ACCENT)
    rect(s, 0, 7.28, 13.333, 0.22, ACCENT_D)
    textbox(s, 1.17, 0.62, 11, 0.4,
            [{"runs": [(num, 14, False, TINT_B2)], "align": PP_ALIGN.CENTER}])
    textbox(s, 1.17, 1.55, 11, 1.0,
            [{"runs": [(big, 40, True, WHITE)], "align": PP_ALIGN.CENTER}])
    rect(s, 5.17, 2.62, 3, 0.02, TINT_B)
    textbox(s, 1.17, 2.82, 11, 0.5,
            [{"runs": [(sub, 16, False, TINT_B)], "align": PP_ALIGN.CENTER}])
    textbox(s, 2.67, 3.9, 8, 2.2, [
        {"runs": [("•  " + p, 13, False, WHITE)],
         "line": 1.4, "space_after": 10, "align": PP_ALIGN.LEFT}
        for p in points])
    footer(s, PAGE)
    return s


# ---- 1 封面 ----
s = new_slide()
rect(s, 0, 0, 13.333, 7.5, ACCENT)
rect(s, 0, 7.28, 13.333, 0.22, ACCENT_D)
textbox(s, 1.17, 0.95, 11, 0.4,
        [{"runs": [("人工智能综合实践课程设计 · 课题编号 2", 13, False, TINT_B2)],
          "align": PP_ALIGN.CENTER}])
textbox(s, 1.17, 2.35, 11, 1.1,
        [{"runs": [("主模型与桌面应用", 40, True, WHITE)], "align": PP_ALIGN.CENTER}])
rect(s, 4.67, 3.62, 4, 0.02, TINT_B)
textbox(s, 1.17, 3.82, 11, 0.45,
        [{"runs": [("基于纹理特征的工业缺陷检测系统 · 模块讲解", 17, False, TINT_B)],
          "align": PP_ALIGN.CENTER}])
textbox(s, 1.17, 4.95, 11, 0.4,
        [{"runs": [("汇报人：陈宇（202305630103）　　指导教师：刘焕华", 14, False, WHITE)],
          "align": PP_ALIGN.CENTER}])
textbox(s, 1.17, 5.5, 11, 0.4,
        [{"runs": [("2026 年 10 月", 12, False, TINT_B2)], "align": PP_ALIGN.CENTER}])
for k, c in enumerate([BLUE_C, ORG_C, RGBColor(0x1B, 0xAF, 0x7A)]):
    rect(s, 5.92 + k * 0.3, 6.85, 0.22, 0.22, c)

# ---- 2 目录 ----
s = new_slide()
header(s, "", "目录")
toc = [
    ("01", "模块一 · 主模型 —— PatchCore 主方案"),
    ("　", "原理与最终配置 · 优化议程 · 批量前向 + GPU 提速 · 聚合消融 · 最终指标"),
    ("02", "模块二 · 桌面应用 —— 质检判定系统"),
    ("　", "界面与流程 · 自动阈值设计 · PatchCore 留一法 · 测试集判定验证"),
    ("03", "总结与展望"),
]
y = 1.6
for num, t in toc:
    textbox(s, 1.2, y, 1.2, 0.5, [{"runs": [(num, 17, True, ACCENT2)]}])
    textbox(s, 2.5, y + 0.06, 10.0, 0.5, [{"runs": [(t, 15, False, INK_C)]}])
    rect(s, 1.2, y + 0.55, 11.0, 0.012, HAIR)
    y += 0.66
footer(s, PAGE)

# ---- 3 模块一分隔 ----
divider("MODULE 1", "模块一 · 主模型",
        "PatchCore 主方案：多层特征记忆库 + kNN 检索",
        ["最终指标：图像级 AUROC 0.9477 / 像素级 0.9240（像素级全项目第一）",
         "推理提速：批量前向 + GPU 分块 GEMM 检索，单张 216.30 → 25.96 ms（8.3×）",
         "聚合消融：三模型全测 max / p95 / p99，确定统一保留 max（Anomalib 协议）"])

# ---- 4 主模型原理与配置 ----
s = content_slide("01", "主模型原理与最终配置", "PatchCore：冻结骨干提取局部特征，靠记忆库检索度量纹理偏离")
steps = [
    ("STEP 1", "特征提取", "冻结 VGG16\nconv3_3 + conv4_3\n多层 patch 特征"),
    ("STEP 2", "记忆库", "220 张正常样本\n≈ 5.5 万 patch 全量入库\n（coreset 1.0 · stride 1）"),
    ("STEP 3", "检索", "k=1 最近邻距离\nGPU 分块 GEMM /\nFAISS（数学等价）"),
    ("STEP 4", "得分", "距离热力图\n(resize + smooth)\n图像级取 max"),
]
x = 0.55
for t, d1, d2 in steps:
    rect(s, x, 1.7, 2.8, 1.95, ALT_C, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    textbox(s, x + 0.25, 1.88, 2.3, 0.35, [{"runs": [(t, 10.5, True, ACCENT2)]}])
    textbox(s, x + 0.25, 2.3, 2.3, 0.4, [{"runs": [(d1, 13, True, INK_C)]}])
    textbox(s, x + 0.25, 2.75, 2.3, 0.9, [{"runs": [(d2, 10, False, SEC_C)], "line": 1.25}])
    if x < 9.5:
        rect(s, x + 2.88, 2.72, 0.22, 0.32, ACCENT2, shape=MSO_SHAPE.RIGHT_ARROW)
    x += 3.08
rect(s, 0.55, 4.1, 12.23, 1.0, TINT_B, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, 0.85, 4.24, 11.7, 0.8, [
    {"runs": [("最终配置（实验五搜索确定）：", 12, True, ACCENT_D),
              ("多层特征融合（conv3_3 + conv4_3）+ k=1 + 全量记忆库（无 coreset 压缩）+ stride 1；"
               "conv3_3 保留细密纹理、conv4_3 提供高层语义。",
               12, False, INK_C)], "line": 1.3}])
rect(s, 0.55, 5.45, 12.23, 1.05, ALT_C, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, 0.85, 5.6, 11.7, 0.85, [
    {"runs": [("指标（metal_nut 测试集 115 张）：", 12, True, ACCENT_D),
              ("图像级 AUROC 0.9477（像素级 0.9240，全项目第一）· 单张推理 25.96 ms（<1s）· "
               "与 FAISS 检索逐位一致，协议不变。", 12, False, INK_C)], "line": 1.3}])
footer(s, PAGE)

# ---- 5 优化议程 ----
s = content_slide("02", "优化议程：依据同学《优化实验报告》逐项验证")
rows = [
    ["议程（同学优化报告）", "我们的验证结论", "去向"],
    ["① 图像级聚合：95 分位替代 max", "三模型全测：95 分位均不如 max（强模型上稀释信号）", "保留 max → 详见 P7"],
    ["② 批量前向 + GPU 推理提速", "已落地：PatchCore 8.3×、PaDiM 6.1×，AUROC 逐位一致", "采纳 → 详见 P6"],
    ["③ PaDiM 配置对齐论文检查", "VGG16 features.22/30 + 550 维随机投影，符合论文设置", "无需改动"],
]
make_table(s, 0.55, 1.75, 12.23, rows, col_w=[3.6, 6.0, 2.6],
           col_aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER], row_h=0.75,
           font_size=11.5, header_size=12,
           fill_cells={(1, 2): TINT_B, (2, 2): TINT_B, (3, 2): ALT_C},
           bold_cells={(1, 2): True, (2, 2): True},
           color_cells={(1, 2): ACCENT_D, (2, 2): ACCENT_D})
rect(s, 0.55, 4.85, 12.23, 1.05, ALT_C, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, 0.85, 5.0, 11.7, 0.85, [
    {"runs": [("结论：", 12, True, ACCENT_D),
              ("同学报告中的优化手段并非全部照搬——95 分位在其弱模型上有效，"
               "在我们的强模型上反而有害；采纳「批量前向 + GPU 提速」并保留 max 聚合协议。",
               12, False, INK_C)], "line": 1.3}])
footer(s, PAGE)

# ---- 6 优化一：批量前向 + GPU 检索 ----
s = content_slide("03", "优化一：批量前向 + GPU 检索（提速 8.3× / 6.1×）")
textbox(s, 0.55, 1.55, 5.6, 0.4, [{"runs": [("技术要点", 14, True, INK_C)]}])
textbox(s, 0.55, 2.0, 5.9, 3.2, bullets([
    ("批量特征提取：一次前向输出整批 patch 特征，省去逐图调度开销",),
    ("GPU 分块 GEMM：d² = |x|² + |y|² − 2xyᵀ，按显存自适应分块，峰值可控",),
    ("k=1 专用路径：行最小值代替 topk，减少排序开销",),
    ("修复隐藏 bug：torch.device(\"cuda\") ≠ 张量的 \"cuda:0\"，比较不相等导致每次调用重建 173MB 记忆库（+37 ms/次），改比较 device.type 后修复",),
    ("AUROC 与 FAISS 检索逐位一致，评估协议不变",),
], size=11.5, gap=9))
textbox(s, 6.85, 1.55, 5.6, 0.4, [{"runs": [("提速结果（单张 / ms，×20 次含 CUDA 同步）", 14, True, INK_C)]}])
make_table(s, 6.85, 2.0, 5.9, [
    ["模型", "优化前", "优化后", "提速"],
    ["PatchCore", "216.30", "25.96", "8.3×"],
    ["PaDiM", "271.94", "44.44", "6.1×"],
    ["CutPaste", "12.43", "12.45", "—"],
], col_w=[2.0, 1.3, 1.3, 1.1], row_h=0.55, font_size=11.5,
   bold_cells={(1, 2): True, (2, 2): True, (1, 3): True, (2, 3): True},
   color_cells={(1, 2): ACCENT_D, (2, 2): ACCENT_D, (1, 3): ORG_C, (2, 3): ORG_C})
rect(s, 6.85, 4.25, 5.9, 1.0, TINT_B, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, 7.05, 4.4, 5.5, 0.8, [
    {"runs": [("PaDiM：", 11.5, True, ACCENT_D),
              ("predict_dataset_fast 绕过 Lightning 引擎直接批量前向 + GPU 马氏距离，与引擎输出逐位一致（diff 0.00000）。",
               11, False, INK_C)], "line": 1.3}])
footer(s, PAGE)

# ---- 7 优化二：聚合方式消融 ----
s = content_slide("04", "优化二：聚合方式消融 —— max vs 95 / 99 分位")
make_table(s, 0.55, 1.75, 12.23, [
    ["模型", "max（现行协议）", "p95", "p99"],
    ["PatchCore", "0.9477", "0.8754 ↓", "0.9418"],
    ["PaDiM", "0.9668", "0.9159 ↓", "0.9717 ↑"],
    ["CutPaste", "0.9003", "0.8661 ↓", "0.8983"],
], col_w=[2.6, 3.2, 3.2, 3.2], row_h=0.6, font_size=13,
   bold_cells={(1, 1): True, (2, 1): True, (3, 1): True},
   color_cells={(1, 1): ACCENT_D, (2, 1): ACCENT_D, (3, 1): ACCENT_D})
textbox(s, 0.55, 4.35, 12.2, 2.4, bullets([
    ("95 分位在三个强模型上普遍下降（PatchCore 0.9477 → 0.8754）——高分位丢弃了最能表征异常的极端 patch，"
     "同学报告中弱模型上的收益不可照搬", True),
    ("PaDiM 的 p99 略升（+0.0049），但为保证三模型评估协议统一（Anomalib 官方 max 协议）不单独采用",),
    ("结论：三模型统一保留 max 聚合，消融结论已写入实验报告与实验 PPT",),
], size=12.5, gap=12))
footer(s, PAGE)

# ---- 8 主模型最终指标 ----
s = content_slide("05", "主模型最终指标（实验一，测试集 115 张）")
make_table(s, 0.55, 1.7, 6.1, [
    ["模型", "图像级", "像素级", "耗时"],
    ["PaDiM", "0.9668", "0.8759", "44.44 ms"],
    ["PatchCore", "0.9477", "0.9240", "25.96 ms"],
    ["CutPaste", "0.9003", "0.6013", "12.45 ms"],
], col_w=[1.7, 1.4, 1.4, 1.5], row_h=0.52, font_size=11.5,
   bold_cells={(1, 0): True, (2, 2): True, (1, 1): True},
   color_cells={(2, 2): ACCENT_D, (1, 1): ACCENT_D})
w = pic(s, RES / "ppt_figs" / "auroc_image.png", 6.95, 1.55, 3.15)
rect(s, 0.55, 4.75, 12.23, 1.55, ALT_C, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, 0.85, 4.92, 11.7, 1.3, bullets([
    ("PatchCore（主方案）：像素级 0.9240 全项目第一、图像级 0.9477 次优，为自实现模型中的最佳路线", True),
    ("PaDiM 图像级最优 0.9668（Anomalib 库对照）；三模型单张推理均 <1s，满足工业实时约束",),
    ("主模型配置已写入 experiment.yaml 并部署至桌面应用推理端",),
], size=12, gap=9))
footer(s, PAGE)

# ---- 9 模块二分隔 ----
divider("MODULE 2", "模块二 · 桌面应用",
        "质检判定系统：自动阈值 + 灵敏度滑杆 + 合格/不合格判定",
        ["交互：上传工件图片 → 三模型切换 → 原图 / 热力图 / 缺陷掩码 / 绿红判定",
         "自动判定阈值由训练集正常样本统计，随权重保存（thresholds.json）",
         "测试集判定验证：PaDiM 检出率 95.7% / 误报率 18.2%，PatchCore 73.1% / 4.5%"])

# ---- 10 桌面应用界面与流程 ----
s = content_slide("06", "桌面应用：界面与流程（Gradio，质检判定）")
steps = [
    ("STEP 1", "启动系统", "python app.py"),
    ("STEP 2", "上传图片", "工件照片\n（png / jpg / bmp）"),
    ("STEP 3", "设置参数", "切换三模型\n调节灵敏度倍数"),
    ("STEP 4", "输出判定", "热力图 · 掩码\n绿「合格」/ 红「不合格」"),
]
x = 0.55
for t, d1, d2 in steps:
    rect(s, x, 1.6, 2.8, 1.7, ALT_C, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    textbox(s, x + 0.25, 1.78, 2.3, 0.35, [{"runs": [(t, 10.5, True, ACCENT2)]}])
    textbox(s, x + 0.25, 2.2, 2.3, 0.4, [{"runs": [(d1, 13, True, INK_C)]}])
    textbox(s, x + 0.25, 2.62, 2.3, 0.6, [{"runs": [(d2, 10, False, SEC_C)], "line": 1.2}])
    if x < 9.5:
        rect(s, x + 2.88, 2.62, 0.22, 0.32, ACCENT2, shape=MSO_SHAPE.RIGHT_ARROW)
    x += 3.08
w = pic(s, DEMO["defect"], 0.55, 3.55, 2.35)
textbox(s, 0.55 + w + 0.25, 3.6, 12.2 - w - 0.3, 2.8, bullets([
    ("异常得分 = 原始异常热力图最大值（Anomalib 协议）",),
    ("缺陷掩码与判定阈值同源：异常得分 ≥ 判定阈值的区域",),
    ("灵敏度倍数滑杆：阈值 = 倍数 × 自动阈值（默认 1.0），调小更严格（少漏检）、调大更宽松（少误报）",),
    ("参考 B2 组螺丝缺陷检测系统（桌面 EXE）的质检交互设计",),
], size=11.5, gap=9))
footer(s, PAGE)

# ---- 11 自动阈值设计 + 留一法 ----
s = content_slide("07", "自动阈值设计：训练集统计 + PatchCore 留一法")
textbox(s, 0.55, 1.5, 5.6, 0.4, [{"runs": [("设计要点", 14, True, INK_C)]}])
textbox(s, 0.55, 1.95, 5.9, 4.3, bullets([
    ("自动阈值 = 训练集正常样本异常得分最大值（220 张，与部署同路径推理）",),
    ("随权重保存 checkpoints/thresholds.json，应用启动自动加载",),
    ("判定阈值 = 灵敏度倍数 × 自动阈值，默认 1.0 = 自动阈值",),
    ("得分 ≥ 阈值 → 红色「不合格」，否则绿色「合格」",),
    ("scripts/compute_auto_threshold.py 一键重算（换数据集 / 重训后更新）",),
], size=11.5, gap=10))
textbox(s, 6.85, 1.5, 5.6, 0.4, [{"runs": [("PatchCore 特例：留一法", 14, True, INK_C)]}])
rect(s, 6.85, 1.95, 5.9, 1.35, RGBColor(0xFC, 0xE8, 0xE1), line=HAIR,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, 7.05, 2.08, 5.5, 1.1, [
    {"runs": [("问题：", 11.5, True, BAD_C),
              ("训练样本自身在记忆库中（自匹配距离 ≈ 0），直接取训练集最大值 → 阈值过低，测试集误报率 100%。",
               11, False, INK_C)], "line": 1.3}])
rect(s, 6.85, 3.5, 5.9, 1.5, TINT_B, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, 7.05, 3.63, 5.5, 1.3, [
    {"runs": [("解决：", 11.5, True, ACCENT_D),
              ("留一法——取第二近邻（k=2）距离：第一近邻是自身，第二近邻即「库外最近真实匹配」，"
               "与部署 k=1 得分同尺度，无需重建记忆库。",
               11, False, INK_C)], "line": 1.3}])
rect(s, 6.85, 5.2, 5.9, 1.05, ALT_C, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
textbox(s, 7.05, 5.33, 5.5, 0.85, [
    {"runs": [("效果：", 11.5, True, GOOD_C),
              ("误报率 100% → 4.5%，检出率 73.1%（自动阈值 1742.6）。", 11.5, True, INK_C)],
     "line": 1.3}])
footer(s, PAGE)

# ---- 12 判定验证 ----
s = content_slide("08", "测试集判定验证（自动阈值 1.0×，115 张）")
make_table(s, 0.55, 1.7, 5.9, [
    ["模型", "缺陷检出率", "误报率"],
    ["PaDiM", "95.7%（89/93）", "18.2%（4/22）"],
    ["PatchCore", "73.1%（68/93）", "4.5%（1/22）"],
    ["CutPaste", "49.5%（46/93）", "0.0%（0/22）"],
], col_w=[1.7, 2.1, 2.1], row_h=0.6, font_size=12,
   bold_cells={(1, 1): True, (1, 2): True, (2, 1): True, (2, 2): True},
   color_cells={(1, 1): GOOD_C, (2, 1): GOOD_C})
w = pic(s, DEMO["good"], 6.8, 1.65, 2.2)
textbox(s, 0.55, 4.35, 12.2, 2.2, bullets([
    ("PaDiM 判定性能最优：检出率 95.7%、误报率 18.2%（马氏距离尺度稳定，最大法阈值有效）", True),
    ("PatchCore 留一法阈值保守（检出 73.1% / 误报 4.5%），产线可调小倍数滑杆提高检出率",),
    ("CutPaste 像素级定位能力弱（PixelAUROC 0.6013），质检判定以 PatchCore / PaDiM 为准",),
    ("灵敏度倍数滑杆在 0.2×~2.0× 间按产线要求权衡漏检与误报",),
], size=12, gap=10))
footer(s, PAGE)

# ---- 13 总结与展望 ----
s = content_slide("09", "总结与展望")
rect(s, 0.55, 1.5, 0.08, 0.32, ACCENT2)
textbox(s, 0.75, 1.45, 5, 0.4, [{"runs": [("总结", 15, True, INK_C)]}])
textbox(s, 0.55, 2.0, 5.9, 4.3, bullets([
    ("主模型：多层融合 + k=1 + 全量记忆库 → 0.9477 / 0.9240（像素级第一）", True),
    ("GPU 提速：PatchCore 216.30 → 25.96 ms（8.3×）、PaDiM 6.1×，指标逐位一致",),
    ("聚合消融：95 分位普遍下降，统一保留 max 协议",),
    ("桌面应用：自动阈值 + 倍数滑杆 + 绿红判定；PaDiM 检出率 95.7%",),
    ("PatchCore 留一法修复自匹配问题，误报率 100% → 4.5%",),
], size=11.5, gap=9))
rect(s, 6.9, 1.5, 0.08, 0.32, ORG_C)
textbox(s, 7.1, 1.45, 5, 0.4, [{"runs": [("展望", 15, True, INK_C)]}])
textbox(s, 6.9, 2.0, 5.9, 4.3, bullets([
    ("PyInstaller 打包独立桌面 EXE（参考 B2 组方案），脱离 Python 环境一键运行",),
    ("扩展到更多类别（screw、cable 等）与更多材质表面",),
    ("模型轻量化与推理加速（ONNX / TensorRT），适配边缘部署",),
], size=12, gap=10))
footer(s, PAGE)

# ---- 14 致谢 ----
s = new_slide()
rect(s, 0, 0, 13.333, 7.5, ACCENT)
rect(s, 0, 7.28, 13.333, 0.22, ACCENT_D)
textbox(s, 1.17, 2.75, 11, 1.2,
        [{"runs": [("谢谢聆听 · 欢迎提问", 40, True, WHITE)], "align": PP_ALIGN.CENTER}])
rect(s, 4.67, 4.15, 4, 0.02, TINT_B)
textbox(s, 1.17, 4.4, 11, 0.5,
        [{"runs": [("基于纹理特征的工业缺陷检测系统 · 主模型与桌面应用", 15, False, TINT_B)],
          "align": PP_ALIGN.CENTER}])
textbox(s, 1.17, 5.15, 11, 0.4,
        [{"runs": [("陈宇（202305630103）· 指导教师：刘焕华", 13, False, WHITE)],
          "align": PP_ALIGN.CENTER}])
footer(s, PAGE)

OUT = ROOT / "基于纹理特征的工业缺陷检测系统_模块讲解.pptx"
prs.save(OUT)
print(f"[ppt] 已保存：{OUT}")
print(f"[ppt] 共 {PAGE} 页")
